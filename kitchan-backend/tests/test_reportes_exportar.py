"""Tests de la exportación de analíticas (Excel y PDF): el caso de uso reutiliza
el dashboard y cada exportador es un adaptador intercambiable."""

from datetime import datetime, timedelta, timezone
from io import BytesIO

import pytest
from openpyxl import load_workbook

from src.kitchan.modules.reportes.application.dashboard_pedidos import (
    ObtenerDashboardPedidosUseCase,
    RangoInvalidoError,
    zona_horaria,
)
from src.kitchan.modules.reportes.application.exportar_analiticas import (
    ExportarAnaliticasPedidosUseCase,
    FormatoExportacionNoSoportado,
)
from src.kitchan.modules.reportes.domain.entities import ItemResumen, PedidoResumen
from src.kitchan.modules.reportes.domain.ports import PedidosLecturaPort
from src.kitchan.modules.reportes.infrastructure.exportador_excel import ExportadorExcel
from src.kitchan.modules.reportes.infrastructure.exportador_pdf import ExportadorPdf

AHORA = datetime(2026, 10, 8, 20, 0, tzinfo=timezone.utc)


class PedidosFalsos(PedidosLecturaPort):
    def __init__(self, pedidos):
        self.pedidos = pedidos

    async def listar_desde(self, restaurante_id, desde):
        return [p for p in self.pedidos if p.fecha_creacion >= desde]


def _dashboard(pedidos):
    return ObtenerDashboardPedidosUseCase(
        pedidos=PedidosFalsos(pedidos), zona=zona_horaria(), ahora=lambda: AHORA
    )


def _pedido(
    origen, total, estado="ENTREGADA", horas_atras=1, cliente="", metodo_pago=None
):
    creado = AHORA - timedelta(hours=horas_atras)
    return PedidoResumen(
        origen=origen,
        total=total,
        estado=estado,
        fecha_creacion=creado,
        fecha_actualizacion=creado + timedelta(minutes=12),
        cliente=cliente,
        metodo_pago=metodo_pago,
        items=[
            ItemResumen(
                nombre="Hamburguesa <Doble> & Papas",
                cantidad=2,
                precio_unitario=total / 2,
            )
        ],
    )


def _caso_de_uso(pedidos):
    return ExportarAnaliticasPedidosUseCase(
        dashboard=_dashboard(pedidos),
        exportadores={"xlsx": ExportadorExcel(), "pdf": ExportadorPdf()},
    )


PEDIDOS = [
    _pedido("RAPPI", 20.0, cliente="Ana Pérez", metodo_pago="TARJETA"),
    _pedido("WHATSAPP", 10.0, cliente="Ana Pérez", metodo_pago="EFECTIVO"),
    _pedido("RAPPI", 5.0, "CANCELADA", cliente="Luis Mora"),
]


@pytest.mark.asyncio
async def test_excel_incluye_las_hojas_y_los_datos_del_periodo():
    contenido, exportador, nombre = await _caso_de_uso(PEDIDOS).ejecutar(
        "r1", "xlsx", "7d"
    )

    libro = load_workbook(BytesIO(contenido))
    assert libro.sheetnames == [
        "Resumen",
        "Por canal",
        "Por día",
        "Por hora",
        "Día de la semana",
        "Productos",
        "Métodos de pago",
        "Clientes",
    ]
    assert libro["Resumen"]["B2"].value == 3  # pedidos totales
    canales = {fila[0].value for fila in libro["Por canal"].iter_rows(min_row=2)}
    assert canales == {"RAPPI", "WHATSAPP"}
    metodos = {
        f[0].value: f[1].value
        for f in libro["Métodos de pago"].iter_rows(min_row=2, max_row=3)
    }
    assert metodos == {"TARJETA": 1, "EFECTIVO": 1}
    indicadores = {
        f[0].value: f[1].value
        for f in libro["Clientes"].iter_rows(min_row=2, max_row=7)
    }
    assert (
        indicadores["Clientes únicos"] == 1
    )  # el pedido cancelado de Luis Mora no cuenta
    assert indicadores["Clientes recurrentes"] == 1  # Ana Pérez pidió dos veces
    resumen = {f[0].value for f in libro["Resumen"].iter_rows(min_row=2)}
    assert {
        "Hora pico",
        "Productos distintos vendidos",
        "Umbral de demora (min)",
    } <= resumen
    assert exportador.extension == "xlsx"
    assert nombre == "kitchan-analiticas-7d-2026-10-08.xlsx"


@pytest.mark.asyncio
async def test_pdf_se_genera_con_datos_con_caracteres_especiales_y_sin_pedidos():
    con_datos, _, nombre = await _caso_de_uso(PEDIDOS).ejecutar("r1", "pdf", "hoy")
    vacio, _, _ = await _caso_de_uso([]).ejecutar("r1", "pdf", "hoy")

    assert con_datos.startswith(b"%PDF")
    assert vacio.startswith(b"%PDF")
    assert nombre == "kitchan-analiticas-hoy-2026-10-08.pdf"


@pytest.mark.asyncio
async def test_exportar_filtra_por_canal():
    contenido, _, nombre = await _caso_de_uso(PEDIDOS).ejecutar(
        "r1", "xlsx", "7d", "rappi"
    )

    libro = load_workbook(BytesIO(contenido))
    canales = {fila[0].value for fila in libro["Por canal"].iter_rows(min_row=2)}
    assert canales == {"RAPPI"}
    assert nombre == "kitchan-analiticas-7d-rappi-2026-10-08.xlsx"


@pytest.mark.asyncio
async def test_formato_o_rango_invalido_se_rechaza():
    caso = _caso_de_uso(PEDIDOS)
    with pytest.raises(FormatoExportacionNoSoportado):
        await caso.ejecutar("r1", "docx")
    with pytest.raises(RangoInvalidoError):
        await caso.ejecutar("r1", "xlsx", "1año")


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "formato,firma,extension",
    [("xlsx", b"PK", "xlsx"), ("pdf", b"%PDF", "pdf")],
)
async def test_endpoint_exportar_descarga_archivo(client, formato, firma, extension):
    from src.kitchan.modules.usuarios.infrastructure.security import JWTTokenGenerator

    token = JWTTokenGenerator().generar_token(
        {
            "sub": "a@a.com",
            "id": "1",
            "rol": "ADMIN",
            "restaurante_id": "11111111-1111-1111-1111-111111111111",
        }
    )
    cabeceras = {"Authorization": f"Bearer {token}"}
    ruta = "/api/v1/reportes/pedidos/analiticas/exportar"

    respuesta = await client.get(
        f"{ruta}?formato={formato}&rango=7d", headers=cabeceras
    )
    assert respuesta.status_code == 200
    assert respuesta.content.startswith(firma)
    assert f".{extension}" in respuesta.headers["content-disposition"]

    assert (
        await client.get(f"{ruta}?formato=docx", headers=cabeceras)
    ).status_code == 422
    assert (
        await client.get(f"{ruta}?formato={formato}&rango=99d", headers=cabeceras)
    ).status_code == 422
    assert (await client.get(f"{ruta}?formato={formato}")).status_code in (401, 403)

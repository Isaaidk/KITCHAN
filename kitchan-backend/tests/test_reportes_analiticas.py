"""Tests del caso de uso de Analíticas (módulo reportes) con un puerto de
lectura falso: el cálculo de métricas ya no depende de la base de datos."""

from datetime import date, datetime

import pytest

from src.kitchan.modules.reportes.application.analiticas_pedidos import (
    ObtenerAnaliticasPedidosUseCase,
)
from src.kitchan.modules.reportes.domain.entities import PedidoResumen
from src.kitchan.modules.reportes.domain.ports import PedidosLecturaPort

HOY = date(2026, 10, 8)


class PedidosFalsos(PedidosLecturaPort):
    def __init__(self, pedidos: list[PedidoResumen]):
        self.pedidos = pedidos
        self.consulta = None

    async def listar_desde(self, restaurante_id, desde):
        self.consulta = (restaurante_id, desde)
        return [p for p in self.pedidos if p.fecha_creacion >= desde]


def _pedido(origen, total, estado, creado, actualizado=None):
    return PedidoResumen(
        origen=origen,
        total=total,
        estado=estado,
        fecha_creacion=creado,
        fecha_actualizacion=actualizado,
    )


@pytest.mark.asyncio
async def test_calcula_metricas_del_dia():
    fuente = PedidosFalsos(
        [
            _pedido(
                "RAPPI",
                10.0,
                "LISTA",
                datetime(2026, 10, 8, 12, 0),
                datetime(2026, 10, 8, 12, 10),
            ),
            _pedido(
                "UBER_EATS",
                20.0,
                "ENTREGADA",
                datetime(2026, 10, 8, 13, 0),
                datetime(2026, 10, 8, 13, 20),
            ),
            _pedido("RAPPI", 30.0, "CANCELADA", datetime(2026, 10, 8, 13, 30)),
            # De ayer: solo cuenta en la comparación por hora.
            _pedido("RAPPI", 99.0, "ENTREGADA", datetime(2026, 10, 7, 12, 0)),
        ]
    )
    use_case = ObtenerAnaliticasPedidosUseCase(pedidos=fuente, hoy=lambda: HOY)

    datos = await use_case.ejecutar("rest-1")

    assert fuente.consulta == ("rest-1", datetime(2026, 10, 7, 0, 0))
    assert datos["pedidos_totales_hoy"] == 3
    assert datos["ticket_promedio"] == 20.0
    assert datos["tiempo_promedio_preparacion_minutos"] == 15.0
    assert datos["pedidos_cancelados_hoy"] == 1
    assert datos["por_canal"] == {"RAPPI": 2, "UBER_EATS": 1}
    assert datos["comparacion_hoy_vs_ayer"][12] == {"hora": 12, "hoy": 1, "ayer": 1}
    assert datos["comparacion_hoy_vs_ayer"][13] == {"hora": 13, "hoy": 2, "ayer": 0}


@pytest.mark.asyncio
async def test_sin_pedidos_retorna_ceros():
    use_case = ObtenerAnaliticasPedidosUseCase(
        pedidos=PedidosFalsos([]), hoy=lambda: HOY
    )

    datos = await use_case.ejecutar("rest-1")

    assert datos["pedidos_totales_hoy"] == 0
    assert datos["ticket_promedio"] == 0.0
    assert datos["tiempo_promedio_preparacion_minutos"] == 0.0
    assert len(datos["comparacion_hoy_vs_ayer"]) == 24


@pytest.mark.asyncio
async def test_endpoint_analiticas_responde_con_datos_reales(client, db_session):
    """Cadena completa: ruta -> caso de uso -> adaptador -> repositorio de
    pedidos (SQLite en tests), sin que reportes toque el modelo ORM."""
    import uuid

    from src.kitchan.modules.pedidos.domain.entities import Pedido, PedidoItem
    from src.kitchan.modules.pedidos.infrastructure.repository import (
        PostgresPedidoRepository,
    )
    from src.kitchan.modules.usuarios.infrastructure.security import (
        JWTTokenGenerator,
    )

    restaurante_id = str(uuid.uuid4())
    await PostgresPedidoRepository(session=db_session).guardar(
        Pedido(
            restaurante_id=restaurante_id,
            origen="RAPPI",
            cliente="Cliente",
            items=[PedidoItem(nombre="Papas", cantidad=1, precio_unitario=5.0)],
            total=5.0,
        )
    )
    token = JWTTokenGenerator().generar_token(
        {"sub": "a@a.com", "id": "1", "rol": "ADMIN", "restaurante_id": restaurante_id}
    )

    respuesta = await client.get(
        "/api/v1/reportes/pedidos/analiticas",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert respuesta.status_code == 200
    datos = respuesta.json()
    assert set(datos) == {
        "pedidos_totales_hoy",
        "ticket_promedio",
        "tiempo_promedio_preparacion_minutos",
        "pedidos_cancelados_hoy",
        "por_canal",
        "comparacion_hoy_vs_ayer",
    }
    assert len(datos["comparacion_hoy_vs_ayer"]) == 24


_DATOS_EXPORT = {
    "pedidos_totales_hoy": 3,
    "ticket_promedio": 20.0,
    "tiempo_promedio_preparacion_minutos": 15.0,
    "pedidos_cancelados_hoy": 1,
    "por_canal": {"RAPPI": 2, "UBER_EATS": 1},
    "comparacion_hoy_vs_ayer": [
        {"hora": h, "hoy": 1 if h == 12 else 0, "ayer": 0} for h in range(24)
    ],
}


def test_exportador_excel_genera_libro_con_tres_hojas():
    from io import BytesIO

    from openpyxl import load_workbook

    from src.kitchan.modules.reportes.infrastructure.exportador_excel import (
        ExportadorExcel,
    )

    libro = load_workbook(BytesIO(ExportadorExcel().exportar(_DATOS_EXPORT, "2026-10-08")))

    assert libro.sheetnames == ["Resumen", "Por canal", "Hoy vs ayer"]
    assert libro["Resumen"]["B5"].value == 3
    assert libro["Por canal"]["A2"].value == "RAPPI"
    assert libro["Por canal"]["B2"].value == 2
    assert libro["Hoy vs ayer"].max_row == 25


def test_exportador_pdf_genera_pdf_valido_tambien_sin_pedidos():
    from src.kitchan.modules.reportes.infrastructure.exportador_pdf import ExportadorPdf

    vacio = {**_DATOS_EXPORT, "pedidos_totales_hoy": 0, "por_canal": {}}
    assert ExportadorPdf().exportar(_DATOS_EXPORT, "2026-10-08").startswith(b"%PDF")
    assert ExportadorPdf().exportar(vacio, "2026-10-08").startswith(b"%PDF")


@pytest.mark.asyncio
async def test_caso_de_uso_exportar_rechaza_formato_desconocido():
    from src.kitchan.modules.reportes.application.exportar_analiticas import (
        ExportarAnaliticasPedidosUseCase,
        FormatoExportacionNoSoportado,
    )

    use_case = ExportarAnaliticasPedidosUseCase(
        analiticas=ObtenerAnaliticasPedidosUseCase(
            pedidos=PedidosFalsos([]), hoy=lambda: HOY
        ),
        exportadores={},
        hoy=lambda: HOY,
    )

    with pytest.raises(FormatoExportacionNoSoportado):
        await use_case.ejecutar("rest-1", "docx")


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "formato,firma,extension",
    [("xlsx", b"PK", "xlsx"), ("pdf", b"%PDF", "pdf")],
)
async def test_endpoint_exportar_descarga_archivo(client, formato, firma, extension):
    from src.kitchan.modules.usuarios.infrastructure.security import (
        JWTTokenGenerator,
    )

    token = JWTTokenGenerator().generar_token(
        {"sub": "a@a.com", "id": "1", "rol": "ADMIN", "restaurante_id": "11111111-1111-1111-1111-111111111111"}
    )
    cabeceras = {"Authorization": f"Bearer {token}"}

    respuesta = await client.get(
        f"/api/v1/reportes/pedidos/analiticas/exportar?formato={formato}",
        headers=cabeceras,
    )

    assert respuesta.status_code == 200
    assert respuesta.content.startswith(firma)
    assert f".{extension}" in respuesta.headers["content-disposition"]

    invalido = await client.get(
        "/api/v1/reportes/pedidos/analiticas/exportar?formato=docx", headers=cabeceras
    )
    assert invalido.status_code == 422

    sin_token = await client.get(
        f"/api/v1/reportes/pedidos/analiticas/exportar?formato={formato}"
    )
    assert sin_token.status_code in (401, 403)

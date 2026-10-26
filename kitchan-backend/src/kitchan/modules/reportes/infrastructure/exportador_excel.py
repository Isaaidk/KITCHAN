from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src.kitchan.modules.reportes.domain.ports import ExportadorAnaliticasPort

_CABECERA = PatternFill("solid", fgColor="1F2937")
_TEXTO_CABECERA = Font(bold=True, color="FFFFFF")
_DIA_SEMANA = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]


def _tabla(hoja, titulos: list[str], filas: list[list], anchos: list[int]) -> None:
    for col, titulo in enumerate(titulos, start=1):
        celda = hoja.cell(row=1, column=col, value=titulo)
        celda.fill = _CABECERA
        celda.font = _TEXTO_CABECERA
        celda.alignment = Alignment(horizontal="center")
    for fila in filas:
        hoja.append(fila)
    for letra, ancho in zip("ABCDEFGHI", anchos):
        hoja.column_dimensions[letra].width = ancho
    hoja.freeze_panes = "A2"


class ExportadorExcel(ExportadorAnaliticasPort):
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    extension = "xlsx"

    def exportar(self, dashboard: dict) -> bytes:
        libro = Workbook()
        actual, previo = dashboard["resumen"], dashboard["periodo_anterior"]

        hoja = libro.active
        hoja.title = "Resumen"
        metricas = [
            ("Pedidos totales", "pedidos_totales"),
            ("Pedidos cancelados", "pedidos_cancelados"),
            ("Tasa de cancelación (%)", "tasa_cancelacion"),
            ("Ventas totales ($)", "ventas_totales"),
            ("Ticket promedio ($)", "ticket_promedio"),
            ("Tiempo promedio de preparación (min)", "tiempo_promedio_preparacion_minutos"),
            ("Unidades por pedido", "unidades_por_pedido"),
        ]
        _tabla(
            hoja,
            ["Métrica", "Periodo", "Periodo anterior"],
            [[nombre, actual[k], previo[k]] for nombre, k in metricas],
            [42, 16, 18],
        )
        hoja["E1"] = "Periodo"
        hoja["F1"] = f"{dashboard['desde']} al {dashboard['hasta']} ({dashboard['rango']})"
        hoja["E2"] = "Canal"
        hoja["F2"] = dashboard["canal"] or "Todos"
        hoja.column_dimensions["E"].width = 10
        hoja.column_dimensions["F"].width = 30

        _tabla(
            libro.create_sheet("Por canal"),
            [
                "Canal",
                "Pedidos",
                "Cancelados",
                "Ventas ($)",
                "Ticket promedio ($)",
                "Tiempo prom. (min)",
                "Demorados (%)",
                "% pedidos",
                "% ventas",
            ],
            [
                [
                    c["canal"],
                    c["pedidos"],
                    c["cancelados"],
                    c["ventas"],
                    c["ticket_promedio"],
                    c["tiempo_promedio_minutos"],
                    c["tasa_demora"],
                    c["porcentaje_pedidos"],
                    c["porcentaje_ventas"],
                ]
                for c in dashboard["por_canal"]
            ],
            [18, 10, 12, 12, 20, 18, 14, 11, 11],
        )
        _tabla(
            libro.create_sheet("Por día"),
            ["Fecha", "Pedidos", "Ventas ($)"],
            [[d["fecha"], d["pedidos"], d["ventas"]] for d in dashboard["por_dia"]],
            [14, 10, 12],
        )
        _tabla(
            libro.create_sheet("Por hora"),
            ["Hora", "Pedidos", "Periodo anterior"],
            [
                [f"{h['hora']:02d}:00", h["pedidos"], h["pedidos_anterior"]]
                for h in dashboard["por_hora"]
            ],
            [10, 10, 18],
        )
        _tabla(
            libro.create_sheet("Día de la semana"),
            ["Día", "Pedidos", "Ventas ($)"],
            [
                [_DIA_SEMANA[d["dia"] % 7], d["pedidos"], d["ventas"]]
                for d in dashboard["por_dia_semana"]
            ],
            [14, 10, 12],
        )
        productos = [
            [seccion, p["nombre"], p["cantidad"], p["ventas"]]
            for seccion, clave in (
                ("Más vendidos", "productos_mas_vendidos"),
                ("Mayor ingreso", "productos_mayor_ingreso"),
                ("Menos vendidos", "productos_menos_vendidos"),
                ("Combos", "combos_mas_vendidos"),
            )
            for p in dashboard[clave]
        ]
        _tabla(
            libro.create_sheet("Productos"),
            ["Ranking", "Producto", "Cantidad", "Ventas ($)"],
            productos,
            [18, 40, 10, 12],
        )

        salida = BytesIO()
        libro.save(salida)
        return salida.getvalue()
